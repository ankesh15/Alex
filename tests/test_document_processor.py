import os
import tempfile
import pytest
import fitz
import docx
from app.services.document_processor import DocumentProcessor


def test_pdf_extraction_multi_page():
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Hello from Page 1 of PDF")

        p2 = doc.new_page()
        p2.insert_text((50, 50), "Welcome to Page 2 of PDF")

        doc.save(tmp_path)
        doc.close()

        pages = DocumentProcessor.extract_text(tmp_path, "pdf")

        assert len(pages) == 2
        assert pages[0]["page_number"] == 1
        assert "Hello from Page 1 of PDF" in pages[0]["text"]
        assert pages[1]["page_number"] == 2
        assert "Welcome to Page 2 of PDF" in pages[1]["text"]
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_docx_extraction():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = docx.Document()
        doc.add_paragraph("First paragraph of sample DOCX document.")
        doc.add_paragraph("Second paragraph with technical details.")
        doc.save(tmp_path)

        results = DocumentProcessor.extract_text(tmp_path, "docx")

        assert len(results) == 1
        assert results[0]["page_number"] is None
        assert "First paragraph" in results[0]["text"]
        assert "Second paragraph" in results[0]["text"]
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_txt_extraction():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as tmp:
        tmp.write("This is a simple text file for document processing verification.")
        tmp_path = tmp.name

    try:
        results = DocumentProcessor.extract_text(tmp_path, "txt")

        assert len(results) == 1
        assert results[0]["page_number"] is None
        assert "simple text file" in results[0]["text"]
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_unsupported_file_extension():
    with tempfile.NamedTemporaryFile(suffix=".exe", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        with pytest.raises(ValueError, match="Unsupported document format"):
            DocumentProcessor.extract_text(tmp_path, "exe")
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
